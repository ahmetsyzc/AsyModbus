from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
D=Document(); sec=D.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=Inches(.7); sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=D.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
D.styles['Normal'].font.size=Pt(11); D.styles['Normal'].paragraph_format.space_after=Pt(7)
D.styles['Normal'].paragraph_format.line_spacing=1.08
D.styles['Title'].font.size=Pt(25)
D.styles['Heading 1'].font.size=Pt(17)
D.styles['Heading 2'].font.size=Pt(13)
code_style=D.styles.add_style('Kod',1); code_style.font.name='Consolas'; code_style.font.size=Pt(9); code_style.paragraph_format.space_after=Pt(0); code_style.paragraph_format.line_spacing=1
foot=sec.footer.paragraphs[0]; foot.alignment=2; foot.add_run('AsyModbus   •   Form bağlantı planı   •   ')
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); foot._p.append(f)
def p(t): D.add_paragraph(t)
def h(t): D.add_heading(t,1)
def sub(t): D.add_heading(t,2)
def bullet(t): D.add_paragraph(t,'List Bullet')
def code(t):
 for line in t.strip().splitlines(): D.add_paragraph(line,'Kod')
def page(): D.add_page_break()
D.add_paragraph('AsyModbus Windows Forms bağlantı planı','Title')
p('Mevcut mimariyle komut kuyruğunu çalıştırma ve birlikte test etme')
p('Hazırlayan Codex • Uygulayan proje sahibi • 6 Ekim 2026')
p('Bu belge, senin hazırlayacağın Windows Forms uygulamasını mevcut BusinessLayer ve veritabanı altyapısına bağlaman için hazırlanmıştır. Formun tasarımı sana ait. Form kendi bağlantı ayarıyla webin kullandığı ASY_MODBUS veritabanına bağlanır; belirli aralıklarla bekleyen komutu kontrol ettirir, röle işlemini ve sonuç kaydını mevcut servis sınıfına yaptırır.')
h('1 Mevcut akış ve sorumluluklar')
p('Webde Durdur → MakineKomutlar tablosunda BEKLIYOR → Form timerı → MakineKomutServisi.KontrolEt() → ISLENIYOR → röle ON → 2 saniye → OFF → BASARILI veya BASARISIZ ve operasyon logu.')
bullet('Form: başlatma, durdurma, timer, durum gösterimi ve güvenli kapanış. Formda makine veya röle yönetim ekranı gerekmiyor.')
bullet('BusinessLayer: komutu sahiplenme, güncel makine ve röle bilgilerini okuma, fiziksel tetikleme ve sonuç kaydı.')
bullet('Veritabanı: mevcut Entity → VeritabaniIslemleri → Stored Procedure zinciri. Form içinde yeni SQL sorguları veya ikinci bir log mekanizması yazma.')
h('2 Codex tarafından yapılanlar')
bullet('MakineKomutlar tablosu, entity sınıfı ve kuyruk prosedürleri eklendi. Aynı makineye ikinci bekleyen veya işlenen talep engellendi; komut sahiplenme atomik yapıldı.')
bullet('Webdeki Durdur işlemi fiziksel röle çağrısı yerine yetki ve makine kontrollerinden sonra kuyruğa talep yazacak şekilde düzenlendi.')
bullet('MakineKomutServisi eklendi. Kendi timerı yok; her KontrolEt çağrısında en fazla bir komut işler. Röle başarısı makinenin durumunu otomatik değiştirmez.')
bullet('Genel loglar VeritabaniIslemleri.Calistir ile; operasyon logları MakinelerLoglar.Ekle ile mevcut altyapıda kaydediliyor. Sonuç ve operasyon logu tek transaction içinde yazılıyor.')
bullet('Eklenen sayısal değer alanı geri alındı. Parametreler.cs önceki hâline döndü. Kontrol aralığını Form yönetiyor; servis süre parametresi okumuyor.')
p('Mevcut doğrulama: çözüm derlendi; geçici veritabanında sahte röleyle 18 servis ve parametre senaryosu geçti. Gerçek röle ve yeni Form bağlantısı henüz test edilmedi.')
page(); h('3 Windows Forms projesini oluştur')
p('1. Visual Studio ile C:\\AsyModbus\\AsyModbus.sln dosyasını aç. Solution üzerine sağ tıkla → Add → New Project seç.')
p('2. Windows Forms App (.NET Framework) şablonunu seç. Örnek proje adı AsyModbus.KomutForm olsun. Hedef frameworkü .NET Framework 4.7.2 seç. Mevcut BusinessLayer da bu sürümü kullanıyor.')
p('3. Form projesindeki References üzerine sağ tıkla → Add Reference → Projects → Solution → BusinessLayer seç. DLL dosyasını elle kopyalamak yerine Project Reference kullan.')
p('4. Formu FrmKomutServisi olarak adlandır. Program.cs içindeki Application.Run satırının bu Formu açtığını kontrol et. Form projesini Set as Startup Project yap.')
p('5. Build Solution çalıştır. MakineKomutServisi mevcut kaynakta global sınıftır; bu sınıf için using BusinessLayer satırı gerekmez. Task için using System.Threading.Tasks kullan.')
h('4 Formun bağlantı ayarını ekle')
p('Form projesine App.config ekle veya mevcut dosyadaki connectionStrings bölümüne aşağıdaki kaydı ekle. Mevcut startup gibi bölümleri silme. Bağlantı adı SqlBaglanti olmalıdır; BusinessLayer bu adı okur.')
code('''<configuration>
  <connectionStrings>
    <add name="SqlBaglanti"
      connectionString="Data Source=DESKTOP-4QHAJ0K\\SQLEXPRESS;Initial Catalog=ASY_MODBUS;Integrated Security=True"
      providerName="System.Data.SqlClient" />
  </connectionStrings>
</configuration>''')
p('Bu bağlantı mevcut yerel SQL instance içindir. Başka bilgisayarda çalıştırırken erişilebilir SQL sunucusunu yaz. Integrated Security kullandığın için Formu çalıştıran Windows hesabının veritabanına erişebilmesi gerekir. Ayrı bir boş veritabanına bağlanma; webin yazdığı kuyruğu okuyabilmesi için aynı ASY_MODBUS kullanılmalıdır.')
p('Çalışma anında okunan ayar Form exe dosyasının config dosyasıdır. BusinessLayer içindeki app.config bu ayarın yerine geçmez. Derleme çıktısında Form exe dosyasının yanında .exe.config bulunduğunu kontrol et.')
h('5 Kontrol aralığını Formda belirle')
p('Form constructor içinde InitializeComponent sonrasında timerKomut.Interval = 5000; yaz. Bu değer beş saniyedir ve kontrol aralığını Form yönetir. Servis çağrıldığında en fazla bir bekleyen komutu işler.')
p('SERVIS_AYARLARI / KOMUT_KONTROL_SURESI_SN kaydı veritabanında kalabilir; servis bu kaydı okumaz. Eksik veya geçersiz olması komut işlemeyi engellemez. Diğer iş parametreleri mevcut mimaride kullanılmaya devam eder. Rölenin iki saniyelik tetikleme süresi kontrol aralığından bağımsızdır.')
page(); h('6 Form kontrollerini ve olayları bağla')
p('Designer üzerinde şu isimlerle kontroller oluştur. İsimleri değiştireceksen aşağıdaki örnekleri de aynı şekilde güncelle.')
bullet('btnBaslat: Başlat düğmesi. Click olayını btnBaslat_Click metoduna bağla.')
bullet('btnDurdur: Durdur düğmesi. Click olayını btnDurdur_Click metoduna bağla.')
bullet('lblDurum: servis durumunu gösteren Label.')
bullet('lstLog: durum mesajlarını gösteren ListBox.')
bullet('timerKomut: System.Windows.Forms.Timer. Enabled = false. Tick olayını timerKomut_Tick metoduna bağla.')
p('FormClosing olayını FrmKomutServisi_FormClosing metoduna bağla. Designer olay bağlantısı yaptıysan aynı olayı ayrıca kodda bağlama; iki kere çalışmasına neden olur. Aşağıdaki parçalar aynı Form sınıfının içine eklenir; Designer.cs içine yazılmaz.')
sub('Alanlar ve constructor')
code('''private readonly MakineKomutServisi servis =
    new MakineKomutServisi();
private bool kontrolAcik;
private bool kapaniyor;
private bool kapanisaIzin;
private Task<bool> kontrolGorevi;

public FrmKomutServisi()
{
    InitializeComponent();
    timerKomut.Enabled = false;
    timerKomut.Interval = 5000;
    servis.LogOlustu += Servis_LogOlustu;
    lblDurum.Text = "Durduruldu";
}''')
sub('Başlat düğmesi')
code('''private void btnBaslat_Click(object sender, EventArgs e)
{
    if (kapaniyor || kontrolAcik ||
        (kontrolGorevi != null && !kontrolGorevi.IsCompleted))
        return;
    kontrolAcik = true;
    lblDurum.Text = "Çalışıyor";
    timerKomut.Start();
}''')
page(); h('7 Timer ve durdurma kodunu ekle')
p('Veritabanı ve röle işlemini Task.Run ile arka planda çalıştır. Timer işlem boyunca durur; Form donmaz ve aynı Formda kontroller üst üste binmez. Her çağrı en fazla bir komut işler; diğer bekleyen komut bir sonraki kontrolde alınır.')
code('''private async void timerKomut_Tick(object sender, EventArgs e)
{
    if (!kontrolAcik || kapaniyor ||
        (kontrolGorevi != null && !kontrolGorevi.IsCompleted))
        return;
    timerKomut.Stop();
    try
    {
        kontrolGorevi = Task.Run(() => servis.KontrolEt());
        bool sonucKaydedildi = await kontrolGorevi;
        if (kapaniyor) return;
        if (servis.AltyapiHatasiVar)
        {
            kontrolAcik = false;
            lblDurum.Text = "Altyapı hatası ile durdu";
            lstLog.Items.Add(servis.SonHata);
        }
        else
        {
            lblDurum.Text = !kontrolAcik ? "Durduruldu" :
                sonucKaydedildi ? "Komut sonucu kaydedildi" :
                "Bekleyen komut kontrol edildi";
        }
    }
    catch (Exception ex)
    {
        kontrolAcik = false;
        if (!kapaniyor)
        {
            lblDurum.Text = "Hata ile durdu";
            lstLog.Items.Add(ex.Message);
        }
    }
    finally
    {
        if (kontrolAcik && !kapaniyor &&
            !IsDisposed && !Disposing)
            timerKomut.Start();
    }
}

private void btnDurdur_Click(object sender, EventArgs e)
{
    kontrolAcik = false;
    timerKomut.Stop();
    lblDurum.Text = "Yeni kontrol durduruldu";
}''')
p('Durdur düğmesi yeni kontrolleri durdurur. Başlamış röle işlemi OFF ve sonuç kaydıyla tamamlanır. KontrolEt true dönerse başarı veya hata sonucu kaydedilmiştir; bu değer tek başına röle başarılı demek değildir. false durumunda AltyapiHatasiVar ile altyapı hatasını boş kuyruktan ayır.')
page(); h('8 Durum mesajlarını ve kapanışı bağla')
sub('Log olayını Form ekranında göster')
p('LogOlustu arka plan threadinden gelebilir. Kontrolleri BeginInvoke ile Form threadinde güncelle. Bu ekran mesajıdır; burada SQL log INSERT yazma. Kalıcı logları mevcut servis ve veritabanı katmanı oluşturur.')
code('''private void Servis_LogOlustu(object sender,
    MakineKomutServisiLogEventArgs e)
{
    if (kapaniyor || IsDisposed || !IsHandleCreated) return;
    try
    {
        BeginInvoke(new Action(() =>
        {
            if (kapaniyor || IsDisposed) return;
            lstLog.Items.Add(e.Tarih.ToString("HH:mm:ss") +
                " " + (e.HataMi ? "HATA " : "") + e.Mesaj);
            if (lstLog.Items.Count > 500)
                lstLog.Items.RemoveAt(0);
            lstLog.TopIndex = lstLog.Items.Count - 1;
        }));
    }
    catch (InvalidOperationException) { }
}''')
sub('Devam eden işlemi bitirerek kapat')
p('Kapanışta timerı durdur, varsa komut görevi bitene kadar UIı engellemeden bekle, ardından Dispose çağır. Task.Wait veya .Result kullanma. Fiziksel işlem devam ederken uygulama sürecini zorla sonlandırma.')
code('''private async void FrmKomutServisi_FormClosing(
    object sender, FormClosingEventArgs e)
{
    if (kapanisaIzin) return;
    e.Cancel = true;
    if (kapaniyor) return;
    kapaniyor = true;
    kontrolAcik = false;
    timerKomut.Stop();
    btnBaslat.Enabled = btnDurdur.Enabled = false;
    lblDurum.Text = "İşlem tamamlanıp kapanıyor";
    try
    {
        if (kontrolGorevi != null) await kontrolGorevi;
    }
    catch (Exception ex)
    {
        lstLog.Items.Add(ex.Message);
    }
    finally
    {
        servis.LogOlustu -= Servis_LogOlustu;
        servis.Dispose();
        kapanisaIzin = true;
        Close();
    }
}''')
page(); h('9 Form tamamlanınca birlikte test edelim')
p('Önce bağlantı ve boş kuyruk testini yapacağız. Bekleyen komut varken gerçek servisle Başlat demek röle tetikleyebilir. Gerçek cihaz testi için uygun makineyi ve zamanı birlikte belirleyeceğiz.')
sub('A Bağlantı ve boş kuyruk')
bullet('Form projesi ve solution derlensin. SqlBaglanti kaydı exe.config içinde bulunsun. Form doğru SQL sunucusuna bağlansın.')
bullet('Kuyrukta BEKLIYOR komut bulunmadığını kontrol edelim. Başlat sonrası Form açık ve kullanılabilir kalsın; röle çağrısı yapılmasın.')
bullet('Formda timerKomut.Interval = 5000 ile beş saniyelik aralık kullanılsın. Durdur sonrası yeni kontrol yapılmasın. Tekrar Başlat çalışsın.')
sub('B Webden tek komut ve sonuç')
bullet('Uygun test makinesinin aktifliği, çalışma durumu, röle IP ve portu ile kanal atamasını kontrol edelim.')
bullet('Webden yetkili kullanıcıyla Durdur talebi oluşturalım. Önce BEKLIYOR kaydı oluştuğunu; Form işleyince BASARILI veya BASARISIZ sonucu yazıldığını doğrulayalım.')
bullet('Gerçek cihazda ON → 2 saniye → OFF akışını gözleyelim. Bir talebin iki kez tetiklenmediğini doğrulayalım.')
bullet('MakinelerLoglar kaydında kaynak SERVIS ve işlem sonucu doğru olsun. Makine durumu sırf röle başarılı diye otomatik değişmesin.')
sub('C Hata ve kapanış')
bullet('Kontrollü bağlantı veya zorunlu iş parametresi hatasında Form timerı dursun ve SonHata gösterilsin. Hata giderildikten sonra Başlat ile tekrar denensin.')
bullet('İşlem sürerken Durdur veya Form kapanışı yeni komut almayı durdursun; başlamış işlem ve sonuç kaydı tamamlansın.')
bullet('Sonuç veya log kaydı başarısızlığında komut ISLENIYOR kalabilir. Bu belirsiz komutu incelemeden BEKLIYOR yapıp tekrar tetiklemeyelim. Bu senaryo önce sahte röleyle test edilsin.')
h('10 Teslim kontrol listesi')
bullet('Windows Forms projesi .NET Framework 4.7.2 ve BusinessLayer Project Reference ile hazır.')
bullet('App.config SqlBaglanti doğru veritabanına bağlı; başlangıç projesi Form.')
bullet('Kontrol isimleri ve beş olay bağlantısı doğru; tek servis nesnesi kullanılıyor.')
bullet('Form içinde bağımsız SQL, ikinci log altyapısı veya makine yönetimi yazılmadı.')
bullet('Form derleniyor; Başlat, Durdur, durum mesajları ve güvenli kapanış hazır.')
p('Hazır olduğunda Form projesini bu solution içinde bırakıp tamamlandığını bildir. Kod bağlantısını inceleyip önce sahte röle ve boş kuyruk kontrollerini, ardından birlikte gerçek cihaz testini yapacağız. Bu belgedeki Form kodu bağlantı örneğidir; senin oluşturacağın proje içinde derlenip doğrulanacaktır.')
Path('docs').mkdir(exist_ok=True); D.save('docs/AsyModbus_Form_Uygulama_Plani.docx')
